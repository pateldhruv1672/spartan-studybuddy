from __future__ import annotations

import ast
import hashlib
import html
import io
import json
import mimetypes
import re
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

from bs4 import BeautifulSoup
from pypdf import PdfReader

from ..db import db
from ..config import settings
from .embeddings import embeddings
from pgvector import Vector
from .telemetry import INDEXED_CHUNKS
from ..graph.facts import extract_facts

CODE_EXTENSIONS={'.py','.js','.jsx','.ts','.tsx','.java','.go','.rs','.rb','.php','.swift','.kt','.scala','.sh','.sql','.r','.lua','.cs','.c','.h','.cpp','.hpp','.vue','.svelte'}
TEXT_EXTENSIONS=CODE_EXTENSIONS|{'.md','.mdx','.txt','.json','.yaml','.yml','.toml','.xml','.html','.htm','.css','.scss','.ipynb','.csv'}
SKIP_DIRS={'.git','node_modules','.venv','venv','dist','build','target','coverage','.next','vendor','__pycache__','.cache'}
MAX_FILE_BYTES=2_000_000

@dataclass
class Chunk:
    content:str; kind:str='text'; symbol:str|None=None; start_line:int|None=None; end_line:int|None=None; topic:str|None=None; metadata:dict[str,Any]|None=None
@dataclass
class Symbol:
    name:str; qualified_name:str; kind:str; start_line:int; end_line:int; signature:str=''; docstring:str=''; metadata:dict[str,Any]|None=None


def _hash(data:bytes)->str: return hashlib.sha256(data).hexdigest()
def _clean_html(text:str)->str:
    soup=BeautifulSoup(text,'lxml')
    for t in soup(['script','style','nav','footer','noscript']): t.decompose()
    return '\n'.join(x.strip() for x in soup.get_text('\n').splitlines() if x.strip())

def extract_text(data:bytes,name:str,mime:str|None=None)->tuple[str,str|None]:
    ext=Path(name).suffix.lower(); language=ext.lstrip('.') if ext in CODE_EXTENSIONS else None
    if ext=='.pdf' or mime=='application/pdf':
        return '\n\n'.join((p.extract_text() or '') for p in PdfReader(io.BytesIO(data)).pages),None
    if ext=='.docx':
        with zipfile.ZipFile(io.BytesIO(data)) as z: root=ET.fromstring(z.read('word/document.xml'))
        return '\n'.join(n.text for n in root.iter() if n.tag.endswith('}t') and n.text),None
    if ext=='.pptx':
        out=[]
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for fn in sorted(x for x in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$',x)):
                root=ET.fromstring(z.read(fn)); txt='\n'.join(n.text for n in root.iter() if n.tag.endswith('}t') and n.text)
                if txt: out.append(f'[{Path(fn).stem}]\n{txt}')
        return '\n\n'.join(out),None
    if ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True); out=[]
        for ws in wb.worksheets:
            out.append(f'[Sheet: {ws.title}]')
            for row in ws.iter_rows(values_only=True): out.append('\t'.join('' if v is None else str(v) for v in row))
        return '\n'.join(out),None
    if ext=='.ipynb':
        nb=json.loads(data.decode('utf-8',errors='replace')); out=[]
        for i,c in enumerate(nb.get('cells',[])): out.append(f"[{c.get('cell_type','cell')} {i}]\n{''.join(c.get('source',[]))}")
        return '\n\n'.join(out),'python'
    text=data.decode('utf-8',errors='replace')
    if ext in {'.html','.htm'} or 'html' in (mime or ''): return _clean_html(text),'html'
    return text,language

def _document_chunks(text:str,max_chars:int=2400,overlap:int=250)->list[Chunk]:
    text=text.replace('\x00',' ').strip(); out=[]
    if not text:return out
    # Preserve headings and paragraph boundaries before falling back to windows.
    blocks=re.split(r'(?m)(?=^#{1,4}\s+|^[A-Z][^\n]{2,80}:\s*$)',text)
    buf=''; start_line=1
    for b in blocks:
        if len(buf)+len(b)<max_chars: buf+=("\n" if buf else '')+b; continue
        if buf.strip(): out.append(Chunk(buf.strip(),topic=(buf.strip().splitlines()[0][:140]),start_line=start_line,end_line=start_line+buf.count('\n'))); start_line+=buf.count('\n')+1
        if len(b)<=max_chars: buf=b; continue
        pos=0
        while pos<len(b):
            end=min(len(b),pos+max_chars); piece=b[pos:end]
            out.append(Chunk(piece.strip(),topic=piece.strip().splitlines()[0][:140] if piece.strip() else '',start_line=start_line,end_line=start_line+piece.count('\n')))
            start_line+=piece.count('\n')+1; pos=max(pos+1,end-overlap)
        buf=''
    if buf.strip(): out.append(Chunk(buf.strip(),topic=buf.strip().splitlines()[0][:140],start_line=start_line,end_line=start_line+buf.count('\n')))
    return out

def _python_chunks(text:str)->tuple[list[Chunk],list[Symbol],list[tuple[str,str,str]]]:
    lines=text.splitlines(); chunks=[]; symbols=[]; edges=[]
    try: tree=ast.parse(text)
    except SyntaxError: return _generic_code_chunks(text,'python')
    module_imports=[]
    for node in ast.walk(tree):
        if isinstance(node,(ast.Import,ast.ImportFrom)):
            names=[a.name for a in node.names]; module_imports.extend(names)
    for node in tree.body:
        if isinstance(node,(ast.FunctionDef,ast.AsyncFunctionDef,ast.ClassDef)):
            st=node.lineno; en=getattr(node,'end_lineno',st); name=node.name; kind='class' if isinstance(node,ast.ClassDef) else 'function'
            src='\n'.join(lines[st-1:en]); sig=lines[st-1].strip() if lines else name; doc=ast.get_docstring(node) or ''
            symbols.append(Symbol(name,name,kind,st,en,sig,doc,{'imports':module_imports[:40]})); chunks.append(Chunk(src,'code',name,st,en,f'{kind} {name}',{'signature':sig}))
            for child in ast.walk(node):
                if isinstance(child,ast.Call):
                    target=''
                    if isinstance(child.func,ast.Name): target=child.func.id
                    elif isinstance(child.func,ast.Attribute): target=child.func.attr
                    if target: edges.append((name,target,'calls'))
            if isinstance(node,ast.ClassDef):
                # Methods are symbols too (spec section 7); their chunk is the enclosing class chunk.
                for m in node.body:
                    if isinstance(m,(ast.FunctionDef,ast.AsyncFunctionDef)):
                        mst=m.lineno; men=getattr(m,'end_lineno',mst); qn=f'{name}.{m.name}'
                        symbols.append(Symbol(m.name,qn,'method',mst,men,lines[mst-1].strip() if lines else m.name,ast.get_docstring(m) or '',{'class':name}))
                        for child in ast.walk(m):
                            if isinstance(child,ast.Call):
                                target=''
                                if isinstance(child.func,ast.Name): target=child.func.id
                                elif isinstance(child.func,ast.Attribute): target=child.func.attr
                                if target: edges.append((qn,target,'calls'))
    # Index top-level glue/imports too.
    covered=set()
    for c in chunks:
        if c.start_line and c.end_line: covered.update(range(c.start_line,c.end_line+1))
    glue=[]; glue_start=None
    for i,line in enumerate(lines,1):
        if i not in covered:
            if glue_start is None: glue_start=i
            glue.append(line)
        elif glue:
            txt='\n'.join(glue).strip()
            if txt: chunks.append(Chunk(txt,'code',None,glue_start,i-1,'module-level code'))
            glue=[]; glue_start=None
    if glue:
        txt='\n'.join(glue).strip()
        if txt: chunks.append(Chunk(txt,'code',None,glue_start,len(lines),'module-level code'))
    return chunks or _generic_code_chunks(text,'python')[0],symbols,edges

def _generic_code_chunks(text:str,language:str)->tuple[list[Chunk],list[Symbol],list[tuple[str,str,str]]]:
    lines=text.splitlines(); patterns=[
        re.compile(r'^\s*(?:export\s+)?(?:async\s+)?(?:function|class|interface|type|def|fn|func|struct|enum)\s+([A-Za-z_$][\w$]*)'),
        re.compile(r'^\s*(?:public|private|protected|static|final|async|override|open|internal|suspend|const|let|var|val|fun|class|interface|struct|enum|func|function|def|fn|type|export|abstract|sealed|data|record|\s)+\s*([A-Za-z_$][\w$]*)\s*\([^;]*\)\s*(?:\{|=>|:)')]
    starts=[]
    for i,line in enumerate(lines,1):
        for p in patterns:
            m=p.search(line)
            if m: starts.append((i,m.group(1))); break
    chunks=[]; symbols=[]
    if starts:
        for idx,(st,name) in enumerate(starts):
            en=(starts[idx+1][0]-1 if idx+1<len(starts) else min(len(lines),st+180)); src='\n'.join(lines[st-1:en]).strip()
            if src:
                chunks.append(Chunk(src,'code',name,st,en,f'symbol {name}',{'language':language})); symbols.append(Symbol(name,name,'symbol',st,en,lines[st-1].strip()))
    else:
        for st in range(1,len(lines)+1,100):
            en=min(len(lines),st+119); src='\n'.join(lines[st-1:en]).strip()
            if src: chunks.append(Chunk(src,'code',None,st,en,f'{language} code'))
    return chunks,symbols,[]

# No token-aware tokenizer library is available to this backend (matching the embedding client, which
# also has no local tokenizer), so — consistent with _document_chunks' existing character-windowed
# fallback below — the bound is character-based: a conservative chars-per-token ratio for source code
# (denser in tokens than prose, due to punctuation/identifiers), with a safety margin under the model's
# actual configured context length for tokenizer variance the chunker cannot see.
_CHARS_PER_TOKEN_SAFETY = 3.0
_TOKEN_SAFETY_MARGIN = 0.85

def max_chunk_chars() -> int:
    return max(500, int(settings.embedding_max_model_len * _TOKEN_SAFETY_MARGIN * _CHARS_PER_TOKEN_SAFETY))

def _split_oversized(c: Chunk, max_chars: int, overlap: int = 200) -> list[Chunk]:
    """Deterministic char-windowed fallback for a single semantic chunk (a class, function, or the
    symbol-boundary span _generic_code_chunks produces) that is itself larger than the embedding model
    can safely accept whole. Preserves kind/metadata; start/end line and symbol/topic are adjusted per
    part so citations still point at real lines and duplicate/near-duplicate parts never overlap fully."""
    text = c.content
    if len(text) <= max_chars:
        return [c]
    base_start = c.start_line or 1
    parts: list[Chunk] = []
    pos = 0
    part_no = 0
    while pos < len(text):
        end = min(len(text), pos + max_chars)
        piece = text[pos:end]
        part_no += 1
        seg_start = base_start + text[:pos].count('\n')
        seg_end = base_start + text[:end].count('\n')
        symbol = f'{c.symbol} (part {part_no})' if c.symbol else c.symbol
        topic = f'{c.topic} (part {part_no})' if c.topic else c.topic
        parts.append(Chunk(piece, c.kind, symbol, seg_start, seg_end, topic, c.metadata))
        if end >= len(text):
            break
        pos = max(pos + 1, end - overlap)  # small overlap so a boundary-crossing reference isn't lost; always progresses
    return parts

def _bounded(chunks: list[Chunk]) -> list[Chunk]:
    limit = max_chunk_chars()
    out: list[Chunk] = []
    for c in chunks:
        out.extend(_split_oversized(c, limit))
    return out

def code_chunks(text:str,language:str)->tuple[list[Chunk],list[Symbol],list[tuple[str,str,str]]]:
    # Applied once here rather than inside each chunker: both the AST-based Python path and the
    # pattern-based generic path build chunks spanning a detected symbol boundary with no size cap of
    # their own (unlike _document_chunks' 2400-char windows), so either can produce an oversized chunk
    # for a large class/function. This is the single choke point both funnel through.
    chunks, symbols, edges = _python_chunks(text) if language in {'py','python'} else _generic_code_chunks(text,language)
    return _bounded(chunks), symbols, edges
def _delete_document(conn, document_id: str) -> None:
    # Child chunks, symbols and edges cascade from indexed_documents in PostgreSQL.
    conn.execute('DELETE FROM indexed_documents WHERE id=?', (document_id,))


def index_bytes(*, project_id: str, source_type: str, source_name: str, data: bytes,
                source_uri: str | None = None, mime_type: str | None = None,
                language: str | None = None, metadata: dict[str, Any] | None = None,
                source_id: str | None = None, document_id: str | None = None) -> dict[str, Any]:
    text, detected = extract_text(data, source_name, mime_type)
    return index_text(
        project_id=project_id, source_type=source_type, source_name=source_name, content=text,
        source_uri=source_uri, mime_type=mime_type, language=language or detected,
        metadata=metadata, source_id=source_id, document_id=document_id, content_hash=_hash(data)
    )


def index_text(*, project_id: str, source_type: str, source_name: str, content: str,
               source_uri: str | None = None, mime_type: str | None = 'text/plain',
               language: str | None = None, metadata: dict[str, Any] | None = None,
               source_id: str | None = None, document_id: str | None = None,
               content_hash: str | None = None) -> dict[str, Any]:
    document_id = document_id or str(uuid.uuid5(uuid.NAMESPACE_URL, f'{project_id}:{source_uri or source_name}'))
    is_code = (Path(source_name).suffix.lower() in CODE_EXTENSIONS) or source_type in {'repository_code', 'vscode_code'}
    if is_code:
        lang = language or Path(source_name).suffix.lstrip('.') or 'code'
        chunks, symbols, edges = code_chunks(content, lang)
    else:
        chunks = _document_chunks(content)
        symbols = []
        edges = []

    vectors = embeddings.embed_batches([c.content for c in chunks])
    facts = extract_facts(content, source_name, language)
    with db() as conn:
        _delete_document(conn, document_id)
        conn.execute(
            'INSERT INTO indexed_documents(id,project_id,source_id,source_type,source_name,source_uri,mime_type,language,content_hash,metadata_json) VALUES(?,?,?,?,?,?,?,?,?,?)',
            (document_id, project_id, source_id, source_type, source_name, source_uri, mime_type, language,
             content_hash or hashlib.sha256(content.encode()).hexdigest(), json.dumps(metadata or {}))
        )
        for i, (c, v) in enumerate(zip(chunks, vectors)):
            if len(v) != settings.embedding_dim:
                raise ValueError(
                    f'Embedding dimension mismatch: got {len(v)}, expected {settings.embedding_dim}. '
                    'Check EMBEDDING_DIM and the embedding model.'
                )
            cid = f'{document_id}:{i}'
            conn.execute(
                'INSERT INTO indexed_chunks(id,document_id,project_id,ordinal,kind,symbol,start_line,end_line,content,topic,metadata_json,embedding) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',
                (cid, document_id, project_id, i, c.kind, c.symbol, c.start_line, c.end_line,
                 c.content, c.topic, json.dumps(c.metadata or {}), Vector(v.tolist()))
            )
        conn.execute('INSERT INTO kg_facts(document_id,project_id,facts_json,partial) VALUES(?,?,?,0) ON CONFLICT(document_id) DO UPDATE SET facts_json=excluded.facts_json,partial=0,extracted_at=CURRENT_TIMESTAMP',
                     (document_id, project_id, json.dumps(facts)))
        for s in symbols:
            sid = str(uuid.uuid5(uuid.NAMESPACE_URL, f'{document_id}:{s.qualified_name}:{s.start_line}'))
            conn.execute(
                'INSERT INTO code_symbols(id,project_id,document_id,name,qualified_name,kind,language,start_line,end_line,signature,docstring,metadata_json) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',
                (sid, project_id, document_id, s.name, s.qualified_name, s.kind, language,
                 s.start_line, s.end_line, s.signature, s.docstring, json.dumps(s.metadata or {}))
            )
        for source, target, kind in edges:
            conn.execute(
                'INSERT INTO code_edges(id,project_id,document_id,source_symbol,target_symbol,edge_type) VALUES(?,?,?,?,?,?)',
                (str(uuid.uuid4()), project_id, document_id, source, target, kind)
            )
    INDEXED_CHUNKS.labels(source_type).inc(len(chunks))
    return {
        'document_id': document_id,
        'chunks': len(chunks),
        'symbols': len(symbols),
        'characters': len(content),
        'storage': 'postgresql+pgvector',
    }


def index_path(project_id: str, path: Path, root: Path | None = None,
               source_id: str | None = None, source_type: str = 'repository_code') -> dict[str, Any]:
    if not path.is_file() or path.stat().st_size > MAX_FILE_BYTES:
        return {'skipped': str(path), 'reason': 'size_or_type'}
    ext = path.suffix.lower()
    if ext not in TEXT_EXTENSIONS | {'.pdf', '.docx', '.pptx', '.xlsx'}:
        return {'skipped': str(path), 'reason': 'unsupported'}
    rel = str(path.relative_to(root)) if root else path.name
    data = path.read_bytes()
    mime, _ = mimetypes.guess_type(path.name)
    return index_bytes(
        project_id=project_id, source_type=source_type, source_name=rel, data=data,
        source_uri=path.as_uri(), mime_type=mime, metadata={'path': rel}, source_id=source_id
    )


def index_repository(project_id: str, root: Path, source_id: str | None = None) -> dict[str, Any]:
    files = []
    totals = {'files': 0, 'chunks': 0, 'symbols': 0, 'skipped': 0, 'failed': 0}
    for p in root.rglob('*'):
        if any(part in SKIP_DIRS for part in p.parts):
            continue
        if not p.is_file():
            continue
        try:
            result = index_path(project_id, p, root, source_id)
        except Exception as exc:
            # One file failing (an unexpected encoding, a parser edge case, ...) must not discard an
            # otherwise-valid repository index; record it and keep going, same as the existing
            # 'skipped' path already does for unsupported/oversized files.
            totals['failed'] += 1
            files.append({'path': str(p.relative_to(root)), 'failed': True, 'error': str(exc)[:300]})
            continue
        if 'skipped' in result:
            totals['skipped'] += 1
            continue
        totals['files'] += 1
        totals['chunks'] += result['chunks']
        totals['symbols'] += result['symbols']
        files.append({'path': str(p.relative_to(root)), **result})
    return {**totals, 'root': str(root), 'documents': files[:200]}


def _fts(project_id: str, query: str, limit: int = 60) -> list[dict[str, Any]]:
    if not query.strip():
        return []
    sql = '''
        SELECT c.id,c.content,c.kind,c.symbol,c.start_line,c.end_line,c.topic,
               d.source_name,d.source_uri,d.language,
               ts_rank_cd(c.search_tsv, websearch_to_tsquery('simple', ?)) AS sparse_score
        FROM indexed_chunks c
        JOIN indexed_documents d ON d.id=c.document_id
        WHERE c.project_id=?
          AND c.search_tsv @@ websearch_to_tsquery('simple', ?)
        ORDER BY sparse_score DESC
        LIMIT ?
    '''
    with db() as conn:
        rows = conn.execute(sql, (query, project_id, query, limit)).fetchall()
    return [dict(r) for r in rows]


def _dense(project_id: str, query: str, limit: int = 60) -> list[dict[str, Any]]:
    q = embeddings.embed([query], query=True)[0]
    if len(q) != settings.embedding_dim:
        raise ValueError(f'Query embedding dimension mismatch: got {len(q)}, expected {settings.embedding_dim}.')
    qv = Vector(q.tolist())
    sql = '''
        SELECT c.id,c.content,c.kind,c.symbol,c.start_line,c.end_line,c.topic,
               d.source_name,d.source_uri,d.language,
               (1 - (c.embedding <=> ?)) AS dense_score
        FROM indexed_chunks c
        JOIN indexed_documents d ON d.id=c.document_id
        WHERE c.project_id=? AND c.embedding IS NOT NULL
        ORDER BY c.embedding <=> ?
        LIMIT ?
    '''
    with db() as conn:
        rows = conn.execute(sql, (qv, project_id, qv, limit)).fetchall()
    return [dict(r) for r in rows]


def hybrid_search(project_id: str, query: str, top_k: int = 10) -> list[dict[str, Any]]:
    sparse = _fts(project_id, query, max(top_k * 6, 40))
    dense = _dense(project_id, query, max(top_k * 6, 40))
    fused: dict[str, dict[str, Any]] = {}
    for label, rows, weight in [('bm25', sparse, 1.0), ('dense', dense, 1.15)]:
        for rank, row in enumerate(rows, 1):
            item = fused.setdefault(row['id'], dict(row))
            item['rrf'] = item.get('rrf', 0) + (weight / (60 + rank))
            item[label + '_rank'] = rank
    out = sorted(fused.values(), key=lambda x: x['rrf'], reverse=True)[:top_k]
    for i, row in enumerate(out, 1):
        line = f"{row.get('start_line') or '?'}-{row.get('end_line') or '?'}"
        row['citation'] = f"[{i}] {row['source_name']}:{line}"
        row['preview'] = row['content'][:700]
    return out


def list_documents(project_id: str, limit: int = 200) -> list[dict[str, Any]]:
    with db() as conn:
        rows = conn.execute(
            'SELECT * FROM indexed_documents WHERE project_id=? ORDER BY indexed_at DESC LIMIT ?',
            (project_id, limit),
        ).fetchall()
    out = []
    for row in rows:
        item = dict(row)
        item['metadata'] = json.loads(item.pop('metadata_json') or '{}')
        out.append(item)
    return out


def project_map(project_id: str) -> dict[str, Any]:
    with db() as conn:
        docs = [dict(r) for r in conn.execute(
            'SELECT source_name,source_type,language FROM indexed_documents WHERE project_id=? ORDER BY source_name',
            (project_id,),
        ).fetchall()]
        symbols = [dict(r) for r in conn.execute(
            'SELECT name,kind,language,start_line,end_line,document_id FROM code_symbols WHERE project_id=? LIMIT 800',
            (project_id,),
        ).fetchall()]
        edges = [dict(r) for r in conn.execute(
            'SELECT source_symbol,target_symbol,edge_type FROM code_edges WHERE project_id=? LIMIT 1000',
            (project_id,),
        ).fetchall()]
    langs: dict[str, int] = {}
    for doc in docs:
        if doc.get('language'):
            langs[doc['language']] = langs.get(doc['language'], 0) + 1
    return {'documents': len(docs), 'files': docs[:300], 'languages': langs, 'symbols': symbols, 'edges': edges}
